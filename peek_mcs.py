import sys, io, collections, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

paths = [
    r'dataset\notxbf_com_excels_f4\bc98-2983-907b\train\bc98-2983-907b_0.xlsx',
    r'dataset\notxbf_com_excels_f4\bc98-2983-907b\valid\bc98-2983-907b_0.xlsx',
    r'dataset\txbf_com_excels_f4\bc98-2983-907b\train\bc98-2983-907b_0.xlsx',
    r'dataset\txbf_com_excels_f4\bc98-2983-907b\valid\bc98-2983-907b_0.xlsx',
]
base = r'c:\Users\11192\Desktop\code\2025B'

for p in paths:
    full = os.path.join(base, p)
    wb = openpyxl.load_workbook(full, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    idx = {h: j for j, h in enumerate(header) if h is not None}
    mj, bj, nj = idx['mcs'], idx['beamforming_en'], idx['noise_floor']
    cnt, bfcnt = collections.Counter(), collections.Counter()
    nfv = []
    first = None
    n = 0
    for row in it:
        n += 1
        cnt[row[mj]] += 1
        bfcnt[row[bj]] += 1
        if len(nfv) < 200:
            nfv.append(float(row[nj]))
        if first is None:
            first = row
    print("=" * 70)
    print(p, "| rows =", n)
    print("  mcs unique (sorted):")
    for k in sorted(cnt, key=lambda x: float(x)):
        print(f"    {k:>8}  count={cnt[k]:>5}")
    print("  beamforming_en:", dict(bfcnt))
    print("  noise_floor: min=%.2f max=%.2f mean=%.2f" % (min(nfv), max(nfv), sum(nfv) / len(nfv)))
    c = first[idx['csi_matrix_r0_c0']]
    s = str(c)
    print("  csi cell type:", type(c).__name__, "len:", len(s))
    print("  csi cell head:", s[:160])
    print("  csi cell tail:", s[-80:])
    wb.close()
